import json
import shutil
import tempfile
from pathlib import Path

import openpyxl
from django.conf import settings
from django.test import TestCase, override_settings

from .models import Consignee

SAMPLE = Path(settings.BASE_DIR) / "sample_data"
TMP = tempfile.mkdtemp()


@override_settings(JOBS_DIR=Path(TMP))
class DashboardFlow(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(TMP, ignore_errors=True)

    def _post(self, **extra):
        with open(SAMPLE / "MASTER_INVOICE.xlsx", "rb") as m, open(SAMPLE / "packing_plan.xlsx", "rb") as p:
            return self.client.post("/", dict(master=m, packing=p, **extra))

    def test_upload_builds_everything(self):
        r = self._post()
        self.assertEqual(r.status_code, 302, getattr(r, "content", b"")[:400])
        job = self.client.get(r.url)
        self.assertContains(job, "SC260458-1")
        # Denmark has no consignee yet -> warnings are shown
        self.assertContains(job, "no consignee on file for DENMARK")
        jid = r.url.strip("/").split("/")[-1]
        res = json.loads((Path(TMP) / jid / "result.json").read_text(encoding="utf-8"))
        self.assertEqual(res["totals"]["qty"], 23)
        self.assertEqual(res["totals"]["amount"], 466.25)
        wb = openpyxl.load_workbook(Path(TMP) / jid / "output" / res["files"]["combined"])
        self.assertEqual(wb.sheetnames, ["INV", "PL", "INV-1", "PL-1"])
        d = self.client.get(f"/job/{jid}/download/{res['files']['zip']}")
        self.assertEqual(d.status_code, 200)
        self.assertEqual(self.client.get(f"/job/{jid}/download/../result.json").status_code, 404)

    def test_consignee_removes_placeholder(self):
        Consignee.objects.create(country="denmark", name="TEST CO", address="Street 1\nCopenhagen", marks="MARKS DK")
        r = self._post()
        jid = r.url.strip("/").split("/")[-1]
        res = json.loads((Path(TMP) / jid / "result.json").read_text(encoding="utf-8"))
        self.assertEqual(res["totals"]["warnings"], 0)
        ws = openpyxl.load_workbook(Path(TMP) / jid / "output" / "SC260458.xlsx").worksheets[0]
        self.assertEqual(ws["E13"].value, "TEST CO")
        self.assertIn("Copenhagen", ws["E14"].value)

    def test_bad_file_shows_error(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        r = self.client.post("/", dict(master=SimpleUploadedFile("x.xlsx", b"not excel")))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Could not build the invoices")


    def test_master_is_built_from_packing_plan(self):
        with open(SAMPLE / "packing_plan.xlsx", "rb") as p:
            r = self.client.post("/", dict(packing=p, invoices="SC260458"))
        self.assertEqual(r.status_code, 302, getattr(r, "content", b"")[:400])
        jid = r.url.strip("/").split("/")[-1]
        res = json.loads((Path(TMP) / jid / "result.json").read_text(encoding="utf-8"))
        self.assertEqual([i["no"] for i in res["invoices"]], ["SC260458", "SC260458-1"])
        self.assertEqual(res["built"]["splits"], 1)
        self.assertEqual(res["totals"]["amount"], 466.25)
        page = self.client.get(r.url)
        self.assertContains(page, "Master invoice built from the packing plan")
        self.assertEqual(self.client.get(f"/job/{jid}/download/MASTER_INVOICE.xlsx").status_code, 200)

    def test_needs_some_file(self):
        r = self.client.post("/", dict())
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Upload the packing plan")

    def test_master_tool_page(self):
        self.assertContains(self.client.get("/master/"), "Packing plan")
        with open(SAMPLE / "packing_plan.xlsx", "rb") as p:
            r = self.client.post("/master/", dict(packing=p, invoices="SC260458"))
        self.assertEqual(r.status_code, 200, r.content[:400])
        self.assertContains(r, "SC260458-1")
        import re
        mid = re.search(r"/master/([0-9a-f]{32})/download/", r.content.decode()).group(1)
        d = self.client.get(f"/master/{mid}/download/")
        self.assertEqual(d.status_code, 200)
        self.assertEqual(self.client.get("/master/xyz/download/").status_code, 404)
