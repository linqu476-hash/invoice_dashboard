#!/usr/bin/env python3
"""
packing_to_master.py  --  Packing plan (.xlsx)  ->  MASTER_INVOICE.xlsx

Run this BEFORE generate_invoices.py.

How the packing plan is read
  * columns are found by their header text (row 1), not by letter, so a moved column does not break it
  * a PO line   = a row with INV.NO. and PO Number
  * a carton    = the PO lines above an orange "carton total" row (CARTON NO. = 1, 2, 3 ...)
                  or, for an invoice with a single carton, above the green "1 OF 1" row
  * every carton after the first one of the same invoice is a split:
        SC260458  (carton 1),  SC260458-1  (carton 2),  SC260458-2  (carton 3) ...
  * N.W  = sum of  (Ordered Quantity x weight)  of the carton's PO lines
    G.W  = N.W + CARTON WEIGHT of that carton            (both are written on the carton's first row)
  * AMOUNT = Ordered Quantity x Price Per Unit + Total Adjustments
  * INV DATE = Customer Request Date (CRD) of the PO line

Usage
  python packing_to_master.py "Packing Plan.xlsx" -o MASTER_INVOICE.xlsx
  python packing_to_master.py "Packing Plan.xlsx" -o MASTER_INVOICE.xlsx --invoices SC260458,SC260460
  python packing_to_master.py "Packing Plan.xlsx" -o MASTER_INVOICE.xlsx --sheet "Packing Plan on 19.10. rev3"
"""
import argparse
import re
import sys
from collections import OrderedDict
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP

import openpyxl
from openpyxl.styles import Font

MASTER_HEADER = ["INV NO", "INV DATE", "PO NO", "CUSTOMER PO NO", "ART NO", "CUS NO", "COUNTRY", "DESCRIPTION", "HTS CODE",
                 "AMOUNT", "QTY", "N.W", "G.W"]

# header text in the packing plan (compared lower-case, trimmed)
WANT = dict(inv="inv.no.", po="po number", cus="gps customer number", country="country", desc="description",
            hts="hts code1:", price="price per unit", adj="total adjustments", unit_w="weight",
            carton_w="carton weight", carton_no="carton no.",
            cpo="customer number", art="article number",
            crd="customer request date (crd)")
CARTON_END_TEXT = re.compile(r"^\s*\d+\s+OF\s+\d+\s*$", re.I)


class PackingError(ValueError):
    pass


def _num(v):
    if v is None or v == "":
        return None
    if isinstance(v, str) and v.startswith("="):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _date(v):
    """CRD cell -> date (a real date stays one; text like 2026-10-19 or 19/10/2026 is parsed; else None)."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        for f in ("%Y-%m-%d", "%d/%m/%Y", "%d.%m.%Y", "%d-%m-%Y", "%B %d, %Y"):
            try:
                return datetime.strptime(v.strip(), f).date()
            except ValueError:
                pass
    return None


def _txt(v):
    return None if v is None or str(v).strip() == "" else str(v).strip()


def _money(x):
    return float(Decimal(str(round(x, 6))).quantize(Decimal("0.01"), ROUND_HALF_UP))


def find_sheet(wb, name=None):
    """The sheet to read: the one asked for, else the first one whose row 1 has INV.NO. and PO Number."""
    if name:
        if name not in wb.sheetnames:
            raise PackingError(f"Sheet '{name}' not found. Sheets: {wb.sheetnames}")
        return wb[name]
    for ws in wb.worksheets:
        head = {str(c.value).strip().lower() for c in ws[1] if c.value}
        if "inv.no." in head and "po number" in head:
            return ws
    raise PackingError("No sheet with the headers 'INV.NO.' and 'PO Number' in row 1.")


def column_map(ws):
    hdr = [str(c.value).strip().lower() if c.value is not None else "" for c in ws[1]]
    optional = {"adj", "carton_w", "carton_no", "unit_w", "cpo", "art", "crd"}
    cols, missing = {}, []
    for key, text in WANT.items():
        if text in hdr:
            cols[key] = hdr.index(text) + 1
        elif key not in optional:
            missing.append(text)
    qtys = [i + 1 for i, h in enumerate(hdr) if h == "ordered quantity"]
    if qtys:
        cols["qty"] = qtys[0]                      # order level quantity
        cols["uqty"] = qtys[-1]                    # size level quantity (used for the weight)
    else:
        missing.append("ordered quantity")
    if missing:
        raise PackingError(f"Packing plan is missing columns: {missing}")
    return cols


def extract(path, sheet=None, only=None):
    """Returns (rows, notes). rows are master rows; notes are warnings (text)."""
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = find_sheet(wb, sheet)
    c = column_map(ws)
    g = lambda r, k: ws.cell(r, c[k]).value if k in c else None
    notes, cartons = [], []
    seen_rows, dropped = 0, {}                      # for a useful error when nothing is found
    for k, label in (("art", "Article Number"), ("cpo", "Customer Number"), ("crd", "Customer Request Date (CRD)")):
        if k not in c:
            notes.append(f"Column '{label}' not found in the packing plan -> { {'art': 'ART NO', 'cpo': 'CUSTOMER PO NO', 'crd': 'INV DATE'}[k] } left blank.")
    cur = None                                      # cartons: list of dict(inv, lines, nw, cw)

    def close(r):
        nonlocal cur
        if cur and cur["lines"]:
            if cur["cw"] is None:
                notes.append(f"{cur['inv']}: no CARTON WEIGHT found above row {r} -> G.W. = N.W.")
            cartons.append(cur)
        cur = None

    for r in range(2, ws.max_row + 1):
        inv, po = g(r, "inv"), g(r, "po")
        if inv or po:
            seen_rows += 1
            if not (inv and po):
                dropped["INV.NO. or PO Number is blank"] = dropped.get("INV.NO. or PO Number is blank", 0) + 1
        if inv and po:                                       # PO line
            inv = str(inv).strip()
            if cur and cur["inv"] != inv:
                notes.append(f"Row {r}: invoice changes to {inv} without a carton total row; "
                             f"carton of {cur['inv']} closed here.")
                close(r)
            cur = cur or dict(inv=inv, lines=[], nw=0.0, cw=None, first_row=r)
            q, uq = _num(g(r, "qty")), _num(g(r, "uqty"))
            price, adj = _num(g(r, "price")), _num(g(r, "adj")) or 0.0
            if q is None or price is None:
                why = "quantity/price is a formula (no saved value)" if any(
                    isinstance(g(r, k), str) and str(g(r, k)).startswith("=") for k in ("qty", "price")) \
                    else "quantity or price is empty"
                dropped[why] = dropped.get(why, 0) + 1
                notes.append(f"Row {r} PO {po}: quantity or price is empty -> row skipped.")
                continue
            uw = _num(g(r, "unit_w"))
            if uw is None:
                notes.append(f"Row {r} PO {po}: unit weight is empty -> counted as 0.")
                uw = 0.0
            cur["nw"] += (uq if uq is not None else q) * uw
            cw = _num(g(r, "carton_w"))
            if cw:
                cur["cw"] = (cur["cw"] or 0.0) + cw
            h = g(r, "hts")
            cur["lines"].append(dict(
                po=int(po) if str(po).isdigit() else po, cus=g(r, "cus"),
                cpo=_txt(g(r, "cpo")), art=_txt(g(r, "art")), crd=_date(g(r, "crd")), country=str(g(r, "country") or "").strip(),
                desc=str(g(r, "desc") or "").strip(),
                hts=int(h) if isinstance(h, (int, float)) else h,
                amount=_money(q * price + adj), qty=int(q)))
            continue
        # not a PO line: carton total (orange) or invoice total (green) rows
        no = g(r, "carton_no")
        is_carton_total = isinstance(no, (int, float)) or (isinstance(no, str) and no.strip().isdigit())
        is_invoice_total = isinstance(no, str) and bool(CARTON_END_TEXT.match(no))
        if is_carton_total or is_invoice_total:
            close(r)                                         # an invoice total row after cartons has nothing open
    if cur and cur["lines"]:
        notes.append(f"Last carton of {cur['inv']} has no closing total row; it was kept.")
        close(ws.max_row)

    # splits: 2nd carton of an invoice -> -1, 3rd -> -2 ...
    seen, rows = {}, []
    for ct in cartons:
        n = seen.get(ct["inv"], 0)
        seen[ct["inv"]] = n + 1
        no = ct["inv"] if n == 0 else f"{ct['inv']}-{n}"
        if only and ct["inv"] not in only and no not in only:
            continue
        nw = round(ct["nw"], 3)
        gw = round(ct["nw"] + (ct["cw"] or 0.0), 3)
        for i, ln in enumerate(ct["lines"]):
            rows.append(dict(inv=no, date=ln["crd"], po=ln["po"], cpo=ln["cpo"], art=ln["art"], cus=ln["cus"], country=ln["country"], desc=ln["desc"],
                             hts=ln["hts"], amount=ln["amount"], qty=ln["qty"],
                             nw=nw if i == 0 else None, gw=gw if i == 0 else None))
    if only:
        found = {r["inv"].split("-")[0] for r in rows} | {r["inv"] for r in rows}
        for want in only:
            if want not in found:
                notes.append(f"Invoice {want} was not found in the packing plan.")
    if not rows:
        detail = "; ".join(f"{n} rows: {k}" for k, n in dropped.items()) or "no row has INV.NO. and PO Number filled"
        tip = "" if sheet else " (no sheet name given -> the first sheet with the headers was used; try typing the sheet name)"
        if only and cartons:
            detail = f"the invoice filter {sorted(only)} matched none of the {len({c['inv'] for c in cartons})} invoices in the sheet"
        raise PackingError(f"No PO lines found in sheet '{ws.title}'{tip}: {detail}.")
    return rows, notes


def write_master(rows, out):
    """Same layout as MASTER_INVOICE.xlsx: header, one block per invoice, a blank row between blocks."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(MASTER_HEADER)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    prev = None
    for r in rows:
        if prev is not None and r["inv"] != prev:
            ws.append([])
        prev = r["inv"]
        ws.append([r["inv"], r["date"], r["po"], r["cpo"], r["art"], r["cus"], r["country"], r["desc"], r["hts"],
                   r["amount"], r["qty"], r["nw"], r["gw"]])
    for row in ws.iter_rows(min_row=2):
        row[11].number_format = row[12].number_format = "0.0##"
        row[9].number_format = "0.00"
        row[1].number_format = "dd-mmm-yyyy"
    for col, w in zip("ABCDEFGHIJKLM", (13, 13, 13, 16, 10, 11, 16, 62, 15, 10, 7, 8, 8)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    wb.save(out)
    return out


def summarize(rows):
    inv = OrderedDict()
    for r in rows:
        d = inv.setdefault(r["inv"], dict(inv=r["inv"], country=r["country"], po=0, qty=0, amount=0.0, nw=None, gw=None))
        d["po"] += 1
        d["qty"] += r["qty"]
        d["amount"] = round(d["amount"] + r["amount"], 2)
        if r["nw"] is not None:
            d["nw"], d["gw"] = r["nw"], r["gw"]
    return list(inv.values())


def main():
    ap = argparse.ArgumentParser(description="Packing plan -> MASTER_INVOICE.xlsx")
    ap.add_argument("packing")
    ap.add_argument("-o", "--out", default="MASTER_INVOICE.xlsx")
    ap.add_argument("--sheet", help="sheet name (default: first sheet that has INV.NO. and PO Number)")
    ap.add_argument("--invoices", help="only these invoices, comma separated, e.g. SC260458,SC260460")
    a = ap.parse_args()
    only = {x.strip() for x in a.invoices.split(",")} if a.invoices else None
    try:
        rows, notes = extract(a.packing, a.sheet, only)
    except PackingError as e:
        sys.exit(f"ERROR: {e}")
    write_master(rows, a.out)
    for s in summarize(rows):
        print(f"{s['inv']:<13}{s['country']:<16}{s['po']:>4} PO  {s['qty']:>4} pcs  {s['amount']:>10.2f} USD  "
              f"N.W {s['nw']}  G.W {s['gw']}")
    print(f"-> {a.out}: {len(rows)} PO lines, {sum(r['qty'] for r in rows)} pcs")
    for n in notes:
        print("WARNING", n)


if __name__ == "__main__":
    main()
