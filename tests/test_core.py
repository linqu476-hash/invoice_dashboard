"""Tests of the service layer with SYNTHETIC data only (no customer data).   python -m unittest discover -s tests -v"""
import datetime
import io
import json
import os
import sqlite3
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from invoicing import extract, shipmarks  # noqa: E402

_TMP = tempfile.mkdtemp()
os.environ["INVOICE_DB_PATH"] = str(Path(_TMP) / "test.sqlite3")
from app_core import runner, store, uploads  # noqa: E402


def fresh_db():
    p = Path(os.environ["INVOICE_DB_PATH"])
    if p.exists():
        p.unlink()


def make_master(country="UNITED KINGDOM", cus="999001", date=datetime.datetime(2026, 10, 19)):
    rows = [dict(inv="SC900001", date=date, po=4500000001 + i, cpo="AA1", art=f"KU{i}", cus=cus, country=country,
                 desc="UNISEX SHORTS 100%POLYESTER(100%RECYCLED) KNITTED", hts=6104630000, amount=20.0 + i, qty=2,
                 nw=1.3 if i == 0 else None, gw=1.85 if i == 0 else None, mode="Air Express") for i in range(2)]
    p = Path(_TMP) / "m.xlsx"
    extract.write_master(rows, p)
    return p.read_bytes()


def make_si(cus, name="Acme AG", street="Platz 1", city="1234 Zug", country="Switzerland"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "overall information"
    for row in (["sold to", cus], [], ["consignee", None, "on invoice", "on customs documents"], ["name", None, "x", name],
                ["street and street number", None, "x", street], ["ZIP code and city", None, "x", city], ["country", None, "x", country], ["notify party"]):
        ws.append(row)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


class StoreTests(unittest.TestCase):
    def setUp(self):
        fresh_db()

    def test_manual_rules(self):
        with self.assertRaises(ValueError):
            store.save_consignee("", "n", "addr")            # customer number required
        with self.assertRaises(ValueError):
            store.save_consignee("35", "n", "")              # address required
        store.save_consignee("35", "Acme", "1 Road")
        with self.assertRaises(ValueError):
            store.save_consignee("000035", "x", "y")         # 000035 = 35
        self.assertIn("35", store.consignee_table())

    def test_si_added_unchanged_replaced(self):
        item = store.read_si_bytes("SI-400447.xlsx", make_si("400447"))
        self.assertEqual(store.save_si_consignee(item)[0], "added")
        self.assertEqual(store.save_si_consignee(item)[0], "unchanged")
        self.assertEqual(store.save_si_consignee(dict(item, name="NEW"))[0], "replaced")
        self.assertEqual(len(store.consignees()), 1)
        self.assertEqual(store.consignees()[0]["country"], "SWITZERLAND")

    def test_si_without_customer_number_is_refused(self):
        with self.assertRaises(shipmarks.MarksError):
            store.save_si_consignee(dict(country="X", name="n", address="a", customer_no=""))

    def test_marks_replace_and_conflicts(self):
        store.save_mark("859015", "ACME\nBERLIN")
        store.save_mark("0859015", "ACME2")                   # same number: replaced, not duplicated
        self.assertEqual(len(store.marks()), 1)
        self.assertEqual(store.marks_table()["859015"], "ACME2")
        with self.assertRaises(ValueError):
            store.save_mark("", "x")

    def test_backup_roundtrip(self):
        store.save_consignee("35", "Acme", "1 Road")
        store.save_mark("35", "MARKS")
        data = json.loads(json.dumps(store.export_data()))
        store.import_data(data, replace=True)
        self.assertEqual((len(store.consignees()), len(store.marks())), (1, 1))
        store.import_data(data)                               # merge: no duplicates
        self.assertEqual((len(store.consignees()), len(store.marks())), (1, 1))
        with self.assertRaises(ValueError):
            store.import_data({"nope": 1})

    def test_import_django_db(self):
        p = Path(_TMP) / "old.sqlite3"
        p.unlink(missing_ok=True)
        c = sqlite3.connect(p)
        c.executescript("CREATE TABLE dashboard_consignee(id INTEGER PRIMARY KEY, country TEXT, customer_no TEXT, name TEXT, address TEXT, marks TEXT);"
                        "CREATE TABLE dashboard_shippingmark(id INTEGER PRIMARY KEY, customer_no TEXT, marks TEXT, note TEXT, source TEXT, updated TEXT);"
                        "INSERT INTO dashboard_consignee VALUES(1,'DENMARK','400447','N','A','');"
                        "INSERT INTO dashboard_shippingmark VALUES(1,'859015','M','','f.xls','2026-10-15');")
        c.commit()
        c.close()
        self.assertEqual(store.import_django_db(p), dict(consignees=1, shippingmarks=1))
        with self.assertRaises(ValueError):
            c = sqlite3.connect(Path(_TMP) / "other.sqlite3")
            c.execute("CREATE TABLE x(a)")
            c.commit()
            c.close()
            store.import_django_db(Path(_TMP) / "other.sqlite3")


class UploadTests(unittest.TestCase):
    def setUp(self):
        fresh_db()

    def test_zip_and_plain_and_bad_files(self):
        z = io.BytesIO()
        with zipfile.ZipFile(z, "w") as zf:
            zf.writestr("a/SI-111111.xlsx", make_si("111111"))
            zf.writestr("__MACOSX/._SI-111111.xlsx", b"x")
            zf.writestr("readme.txt", b"hi")
        res = uploads.save_si_uploads([("pack.zip", z.getvalue()), ("SI-222222.xlsx", make_si("222222")),
                                       ("notes.pdf", b"x"), ("bad.xlsx", b"not excel"), ("empty.zip", b"nope")])
        self.assertEqual((res["added"], res["failed"], res["total"]), (2, 3, 5))
        self.assertEqual(len(store.consignees()), 2)


class PipelineTests(unittest.TestCase):
    def setUp(self):
        fresh_db()

    def test_template_loader_disables_unused_external_link_preservation(self):
        with patch("invoicing.engine.openpyxl.load_workbook") as load:
            runner.engine._load_template_workbook("template.xlsx")
        load.assert_called_once_with("template.xlsx", keep_links=False)

    def test_cut_rows_shifts_formulas_without_scanning_full_sheet_grid(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A20"] = "=SUM(A1:A18)"
        # openpyxl itself uses iter_rows() once while deleting rows. The
        # optimized formula pass must not add a second full-grid scan.
        original = ws.iter_rows
        with patch.object(ws, "iter_rows", wraps=original) as iter_rows:
            runner.engine.cut_rows(wb)
        self.assertEqual(iter_rows.call_count, 1)
        self.assertEqual(ws["A18"].value, "=SUM(A1:A16)")

    def test_build_missing_then_fixable_then_rebuild(self):
        run = runner.build_invoices(make_master(country="GERMANY", cus="999001"), "m.xlsx")
        r = run["result"]
        self.assertEqual(r["totals"]["invoices"], 1)
        self.assertTrue(run["report_pdf"].startswith(b"%PDF"))
        self.assertIn(r["files"]["zip"], run["files"])
        self.assertIsNone(r["files"]["combined"])  # default build skips the expensive combined workbook
        with zipfile.ZipFile(io.BytesIO(run["files"][r["files"]["zip"]])) as zf:
            self.assertNotIn("SC900001_INV-PL.xlsx", zf.namelist())
            self.assertTrue(all(info.compress_type == zipfile.ZIP_STORED for info in zf.infolist()))
        self.assertEqual([m["customer_no"] for m in runner.missing_consignees(r)], ["999001"])
        self.assertEqual(runner.fixable(r), [])
        store.save_consignee("999001", "Test Co", "1 Street\nBerlin")
        store.save_mark("999001", "TEST MARKS")
        self.assertEqual(runner.fixable(r), ["SC900001"])
        run2 = runner.rebuild(run)
        self.assertTrue(run2["result"]["invoices"][0]["consignee_ok"])
        self.assertEqual(runner.missing_consignees(run2["result"]), [])


    def test_multiple_invoice_families_are_saved_and_packaged(self):
        wb = openpyxl.load_workbook(io.BytesIO(make_master(country="GERMANY", cus="999001")))
        headers = {str(c.value).strip().upper(): c.column for c in wb.active[1] if c.value is not None}
        wb.active.cell(3, headers["INV NO"], "SC900002")
        data = io.BytesIO()
        wb.save(data)
        run = runner.build_invoices(data.getvalue(), "two-families.xlsx")
        items = set(run["result"]["files"]["items"])
        self.assertEqual(items, {"SC900001.xlsx", "SC900002.xlsx"})
        with zipfile.ZipFile(io.BytesIO(run["files"][run["result"]["files"]["zip"]])) as zf:
            self.assertEqual(set(zf.namelist()), items)
            for name in items:
                self.assertTrue(zf.read(name).startswith(b"PK"))

    def test_invoice_grouping_is_precomputed_once_per_invoice(self):
        # Grouping is shared by INV.BKK, invoice sheets, optional combined output,
        # and the result preview. It should not be repeated by each consumer.
        from invoicing import engine
        original = engine.group_rows
        with patch.object(engine, "group_rows", wraps=original) as grouped:
            run = runner.build_invoices(make_master(country="GERMANY", cus="999001"), "m.xlsx")
        self.assertEqual(run["result"]["totals"]["invoices"], 1)
        self.assertEqual(grouped.call_count, 1)
        self.assertEqual(run["result"]["invoices"][0]["n_grp"], 1)

    def test_combined_workbook_is_optional(self):
        data = make_master(country="GERMANY", cus="999001")
        run = runner.build_invoices(data, "m.xlsx", combine=True)
        combined = run["result"]["files"]["combined"]
        self.assertEqual(combined, "SC900001_INV-PL.xlsx")
        self.assertNotIn(combined, run["files"])  # XLSX bytes are extracted from ZIP only when downloaded
        self.assertEqual(set(run["files"]), {run["result"]["files"]["zip"]})
        self.assertTrue(runner.output_file_bytes(run, combined).startswith(b"PK"))
        with zipfile.ZipFile(io.BytesIO(run["files"][run["result"]["files"]["zip"]])) as zf:
            self.assertIn(combined, zf.namelist())
        self.assertEqual(run["result"]["totals"]["invoices"], 1)

    def test_output_file_bytes_rejects_unlisted_or_path_names(self):
        run = runner.build_invoices(make_master(country="GERMANY", cus="999001"), "m.xlsx")
        with self.assertRaises(ValueError):
            runner.output_file_bytes(run, "../secrets.toml")
        with self.assertRaises(ValueError):
            runner.output_file_bytes(run, "not-in-this-build.xlsx")
        filename = run["result"]["files"]["items"][0]
        self.assertTrue(runner.output_file_bytes(run, filename).startswith(b"PK"))

    def test_invalid_master_lists_problems_and_builds_nothing(self):
        wb = openpyxl.Workbook()
        wb.active.append(["INV NO", "PO NO"])                 # required columns missing
        b = io.BytesIO()
        wb.save(b)
        with self.assertRaises(runner.InvalidMaster) as cm:
            runner.build_invoices(b.getvalue(), "bad.xlsx")
        self.assertTrue(cm.exception.problems)

    def test_standard_packing_plan_extracts_in_read_only_path(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["INV.NO.", "PO Number", "Gps Customer Number", "COUNTRY", "DESCRIPTION", "HTS CODE1:",
                   "Price Per Unit", "Total Adjustments", "weight", "CARTON WEIGHT", "CARTON NO.",
                   "Customer Number", "Article Number", "Customer Request Date (CRD)", "Shipment Method",
                   "Ordered Quantity", "Ordered Quantity"])
        ws.append(["SC900001", "PO-1", "999001", "UNITED KINGDOM", "SHORTS KNITTED", 6104630000,
                   10, 0, 0.5, 1, None, "CPO-1", "ART-1", datetime.datetime(2026, 10, 19),
                   "Air Express", 10, 10])
        ws.append([None, None, None, None, None, None, None, None, None, None, 1])
        p = Path(_TMP) / "standard_packing_stream.xlsx"
        wb.save(p)
        rows, notes = extract.extract(p)
        self.assertEqual(len(rows), 1)
        self.assertEqual((rows[0]["inv"], rows[0]["qty"], rows[0]["amount"]), ("SC900001", 10, 100.0))
        self.assertEqual((rows[0]["nw"], rows[0]["gw"]), (5.0, 6.0))
        self.assertEqual(notes, [])

    def test_read_packing_streaming_columns(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(2, 7, datetime.datetime(2026, 10, 19))  # CRD
        ws.cell(2, 9, "PO-1")
        ws.cell(2, 20, "CUS-1")
        ws.cell(2, 29, "ART-1")
        ws.cell(2, 37, 12)
        ws.cell(2, 38, 2.5)
        ws.cell(2, 41, 0.25)
        p = Path(_TMP) / "packing_readonly.xlsx"
        wb.save(p)
        pack, crd = __import__("invoicing.engine", fromlist=["read_packing"]).read_packing(p)
        self.assertEqual(pack["PO-1"], dict(art="ART-1", cus="CUS-1", price=2.5, qty=12, adj=0.25))
        self.assertEqual(crd, datetime.datetime(2026, 10, 19))

    def test_not_an_excel_file(self):
        with self.assertRaises(runner.InvalidMaster):
            runner.build_invoices(b"a,b,c\n1,2,3\n", "renamed.xlsx")

    def test_size_limit(self):
        with self.assertRaises(ValueError):
            runner.build_invoices(b"x" * (runner.MAX_BYTES + 1), "huge.xlsx")


class FastMergeTests(unittest.TestCase):
    """engine._merge / _unmerge / _copy_sheet only skip work: the sheet must be the same as with the openpyxl methods
    (this also catches an openpyxl upgrade that changes the internals they rely on)."""

    @staticmethod
    def _state(ws):
        cells = []
        for (r, c), x in sorted(ws._cells.items()):
            b = x.border
            cells.append((r, c, type(x).__name__, x.value, b.left.style, b.right.style, b.top.style, b.bottom.style, x.protection.locked))
        return cells, sorted(str(m) for m in ws.merged_cells.ranges)

    def test_same_sheet_as_openpyxl(self):
        from invoicing import engine
        slow, fast = (openpyxl.load_workbook(engine_template()) for _ in range(2))
        s_ws = slow.copy_worksheet(slow.active)
        f_ws = engine._copy_sheet(fast, fast.active)
        self.assertEqual(self._state(s_ws), self._state(f_ws))
        for ws, fast_mode in ((s_ws, False), (f_ws, True)):
            old = [m for m in ws.merged_cells.ranges if m.min_row >= 25]
            if fast_mode:
                engine._unmerge(ws, old)
                engine._merge(ws, 30, 30, 3, 18)
                engine._merge(ws, 31, 33, 3, 10)
            else:
                for m in old:
                    ws.unmerge_cells(str(m))
                ws.merge_cells(start_row=30, end_row=30, start_column=3, end_column=18)
                ws.merge_cells(start_row=31, end_row=33, start_column=3, end_column=10)
        self.assertEqual(self._state(s_ws), self._state(f_ws))


def engine_template():
    from app_core.runner import TEMPLATE_XLSX
    return TEMPLATE_XLSX


if __name__ == "__main__":
    unittest.main()
